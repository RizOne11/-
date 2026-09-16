from __future__ import annotations

import argparse
import asyncio
from collections import Counter
from typing import Any

from openpyxl import Workbook, load_workbook

from .models import ProductMission, Verdict
from .scouts.catalog import HotlineScout, PromScout
from .scouts.epicentr import EpicentrScout
from .scouts.web_shops import WebShopsScout


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def read_missions(path: str) -> list[ProductMission]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    aliases = {"article": {"артикул", "article", "sku", "код"}, "name": {"назва", "название", "name", "товар", "product"}, "brand": {"бренд", "brand", "виробник"}, "model": {"модель", "model"}}
    header_row = 0
    cols: dict[str, int] = {}
    for r in range(1, min(ws.max_row, 30) + 1):
        vals = {_text(ws.cell(r, c).value).lower(): c for c in range(1, ws.max_column + 1)}
        found = {key: next((vals[n] for n in names if n in vals), 0) for key, names in aliases.items()}
        if found["name"] or (found["article"] and sum(bool(x) for x in found.values()) >= 2):
            header_row, cols = r, found
            break
    if not header_row:
        raise ValueError("Need a Name/Назва column; Article, Brand and Model are optional")
    missions = []
    for r in range(header_row + 1, ws.max_row + 1):
        data = {key: _text(ws.cell(r, col).value) if col else "" for key, col in cols.items()}
        if not any(data.values()):
            continue
        source = {k: data[k] for k in ("name", "brand", "model") if data[k]}
        missions.append(ProductMission(article=data["article"] or f"ROW-{r}", source_data=source))
    return missions


def scouts():
    return [EpicentrScout(timeout=15, max_candidates_per_query=12), PromScout(timeout=15, max_candidates_per_query=12), HotlineScout(timeout=15, max_candidates_per_query=12), WebShopsScout(timeout=15, max_candidates_per_query=20, max_per_domain=2)]


async def scan(mission: ProductMission) -> list[dict[str, Any]]:
    rows = []
    for scout in scouts():
        report = await scout.scan(mission)
        for item in report.offers:
            if item.verdict != Verdict.PASS:
                continue
            offer = item.offer
            rows.append({"article": mission.article, "name": mission.source_data.get("name", ""), "brand": mission.source_data.get("brand", ""), "model": mission.source_data.get("model", ""), "source": scout.marketplace.value, "price": float(offer.price) if offer.price is not None else None, "currency": offer.currency, "availability": offer.availability or "", "found_title": offer.title, "url": str(offer.url), "match": round(item.score, 3), "domain": offer.attributes.get("source_domain", ""), "health": report.health.value})
    return rows


def save(rows: list[dict[str, Any]], output: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Offers"
    ws.append(["Артикул", "Назва", "Бренд", "Модель", "Джерело", "Ціна", "Валюта", "Наявність", "Знайдена назва", "URL", "Match", "Домен", "Health"])
    for x in rows:
        ws.append([x["article"], x["name"], x["brand"], x["model"], x["source"], x["price"], x["currency"], x["availability"], x["found_title"], x["url"], x["match"], x["domain"], x["health"]])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]: cell.font = cell.font.copy(bold=True)
    for col, width in {"A":18,"B":42,"C":16,"D":20,"E":14,"F":14,"G":10,"H":18,"I":55,"J":55,"K":10,"L":28,"M":18}.items(): ws.column_dimensions[col].width = width
    summary = wb.create_sheet("Price summary")
    summary.append(["Артикул", "Назва", "Джерело", "Ціна", "Карток"])
    groups = Counter((x["article"], x["name"], x["source"], x["price"]) for x in rows if x["price"] is not None)
    for key, count in sorted(groups.items(), key=lambda z: (z[0][0], z[0][2], z[0][3])): summary.append([*key, count])
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    for cell in summary[1]: cell.font = cell.font.copy(bold=True)
    for col, width in {"A":18,"B":42,"C":16,"D":14,"E":12}.items(): summary.column_dimensions[col].width = width
    wb.save(output)


async def run(input_path: str, output_path: str) -> None:
    missions = read_missions(input_path)
    if not missions:
        raise ValueError("No products found")
    rows = []
    for i, mission in enumerate(missions, 1):
        print(f"[{i}/{len(missions)}] {mission.article}: {mission.source_data.get('name','')}", flush=True)
        rows.extend(await scan(mission))
    save(rows, output_path)
    print(f"READY: {output_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="PUMA production Excel scanner")
    parser.add_argument("input")
    parser.add_argument("-o", "--output", default="PUMA_result.xlsx")
    args = parser.parse_args()
    asyncio.run(run(args.input, args.output))


if __name__ == "__main__":
    main()
